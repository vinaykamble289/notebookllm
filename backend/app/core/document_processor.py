import os
from pathlib import Path
from typing import List
import fitz  # PyMuPDF
from docx import Document as DocxDocument
import pandas as pd

from langchain.schema import Document


class DocumentProcessor:
    """Process various document types and extract text"""
    
    @staticmethod
    def process_pdf(file_path: str) -> List[Document]:
        """Process PDF files using PyMuPDF"""
        documents = []
        
        try:
            pdf_document = fitz.open(file_path)
            total_pages = len(pdf_document)
            
            if total_pages == 0:
                pdf_document.close()
                raise ValueError("PDF file has no pages")
            
            for page_num in range(total_pages):
                page = pdf_document[page_num]
                text = page.get_text()
                
                if text.strip():
                    documents.append(Document(
                        page_content=text,
                        metadata={
                            "source": os.path.basename(file_path),
                            "page": page_num + 1,
                            "file_type": "pdf"
                        }
                    ))
            
            pdf_document.close()
            
            if not documents:
                raise ValueError(
                    f"PDF has {total_pages} pages but no extractable text. "
                    "The PDF may contain only images or be password-protected."
                )
            
            print(f"✅ Processed PDF: {len(documents)} pages with text (out of {total_pages} total pages)")
            
        except Exception as e:
            print(f"❌ Error processing PDF: {e}")
            raise
        
        return documents
    
    @staticmethod
    def process_docx(file_path: str) -> List[Document]:
        """Process DOCX files"""
        documents = []
        
        try:
            doc = DocxDocument(file_path)
            
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            
            text = "\n".join(full_text)
            
            if not text.strip():
                raise ValueError("DOCX file contains no extractable text")
            
            documents.append(Document(
                page_content=text,
                metadata={
                    "source": os.path.basename(file_path),
                    "file_type": "docx"
                }
            ))
            
            print(f"✅ Processed DOCX: {len(documents)} documents")
            
        except Exception as e:
            print(f"❌ Error processing DOCX: {e}")
            raise
        
        return documents
    
    @staticmethod
    def process_txt(file_path: str) -> List[Document]:
        """Process TXT files"""
        documents = []
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            
            if not text.strip():
                raise ValueError("TXT file is empty")
            
            documents.append(Document(
                page_content=text,
                metadata={
                    "source": os.path.basename(file_path),
                    "file_type": "txt"
                }
            ))
            
            print(f"✅ Processed TXT: {len(documents)} documents")
            
        except Exception as e:
            print(f"❌ Error processing TXT: {e}")
            raise
        
        return documents
    
    @staticmethod
    def process_csv(file_path: str) -> List[Document]:
        """Process CSV files"""
        documents = []
        
        try:
            df = pd.read_csv(file_path)
            
            # Convert to text representation
            text = df.to_string()
            
            if text.strip():
                documents.append(Document(
                    page_content=text,
                    metadata={
                        "source": os.path.basename(file_path),
                        "file_type": "csv",
                        "rows": len(df),
                        "columns": len(df.columns)
                    }
                ))
            
            print(f"✅ Processed CSV: {len(documents)} documents")
            
        except Exception as e:
            print(f"❌ Error processing CSV: {e}")
            raise
        
        return documents
    
    @staticmethod
    def process_xlsx(file_path: str) -> List[Document]:
        """Process XLSX files"""
        documents = []
        
        try:
            df = pd.read_excel(file_path)
            
            # Convert to text representation
            text = df.to_string()
            
            if text.strip():
                documents.append(Document(
                    page_content=text,
                    metadata={
                        "source": os.path.basename(file_path),
                        "file_type": "xlsx",
                        "rows": len(df),
                        "columns": len(df.columns)
                    }
                ))
            
            print(f"✅ Processed XLSX: {len(documents)} documents")
            
        except Exception as e:
            print(f"❌ Error processing XLSX: {e}")
            raise
        
        return documents
    
    @classmethod
    def process_file(cls, file_path: str) -> List[Document]:
        """Process file based on extension"""
        file_ext = Path(file_path).suffix.lower()
        
        processors = {
            '.pdf': cls.process_pdf,
            '.docx': cls.process_docx,
            '.doc': cls.process_docx,
            '.txt': cls.process_txt,
            '.csv': cls.process_csv,
            '.xlsx': cls.process_xlsx,
            '.xls': cls.process_xlsx,
        }
        
        processor = processors.get(file_ext)
        if not processor:
            raise ValueError(f"Unsupported file type: {file_ext}")
        
        return processor(file_path)


document_processor = DocumentProcessor()
